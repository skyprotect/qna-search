import os
import sys
import re
import json
import unicodedata
import openpyxl

def remove_accents(input_str: str) -> str:
    """Loại bỏ dấu tiếng Việt để phục vụ tìm kiếm không dấu."""
    if not input_str:
        return ""
    # Normalize unicode to NFD and strip combining characters
    nfd_str = unicodedata.normalize('NFD', input_str)
    no_accent = ''.join(c for c in nfd_str if unicodedata.category(c) != 'Mn')
    # Thay thế đ/Đ
    no_accent = no_accent.replace('đ', 'd').replace('Đ', 'D')
    return no_accent.lower()

def clean_option_text(key: str, text: str) -> str:
    """Làm sạch tiền tố nếu text đã bị lặp lại key (ví dụ 'D. D. Bản lĩnh...' hoặc 'A. Bản lĩnh...')."""
    if not text:
        return ""
    text = text.strip()
    # Loại bỏ tiền tố trùng lặp như 'A.', 'A -', 'A:', 'A. '
    pattern = rf'^{re.escape(key)}[\.\:\-\s]+(.*)$'
    match = re.match(pattern, text, re.IGNORECASE)
    if match:
        return match.group(1).strip()
    return text

def parse_excel(filepath: str):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    
    questions = []
    current_q = None
    
    for r in range(2, sheet.max_row + 1):
        tt_val = sheet.cell(r, 1).value
        q_val = sheet.cell(r, 2).value
        opt_key_val = sheet.cell(r, 3).value
        opt_text_val = sheet.cell(r, 4).value
        ans_mark_val = sheet.cell(r, 5).value
        
        # Nếu có nội dung câu hỏi mới
        if q_val is not None and str(q_val).strip():
            if current_q:
                questions.append(finalize_question(current_q))
            
            # Khởi tạo câu hỏi mới
            q_clean = str(q_val).strip()
            tt = int(tt_val) if tt_val is not None and str(tt_val).strip().isdigit() else (len(questions) + 1)
            
            current_q = {
                'id': tt,
                'q': q_clean,
                'raw_options': [],
                'correct_key': None,
                'correct_text': None
            }
        
        # Thêm lựa chọn đáp án
        if current_q and (opt_key_val is not None or opt_text_val is not None):
            key = str(opt_key_val).strip().upper() if opt_key_val is not None else ''
            text = str(opt_text_val).strip() if opt_text_val is not None else ''
            cleaned_text = clean_option_text(key, text)
            
            is_correct = False
            if ans_mark_val is not None:
                mark_str = str(ans_mark_val).strip().upper()
                if mark_str in ['X', '1', 'TRUE', 'ĐÚNG', 'DUNG', 'V']:
                    is_correct = True
            
            current_q['raw_options'].append({
                'key': key,
                'text': cleaned_text,
                'correct': is_correct
            })
            
            if is_correct:
                current_q['correct_key'] = key
                current_q['correct_text'] = cleaned_text
    
    if current_q:
        questions.append(finalize_question(current_q))
        
    return questions

def finalize_question(q_obj):
    qid = q_obj['id']
    q_text = q_obj['q']
    correct_key = q_obj['correct_key'] or ''
    correct_text = q_obj['correct_text'] or ''
    
    # Chuỗi đáp án đầy đủ hiển thị (ví dụ: 'C. Từ ngày 20/01/2026...')
    if correct_key and correct_text:
        full_a = f"{correct_key}. {correct_text}"
    elif correct_text:
        full_a = correct_text
    else:
        full_a = ""
        
    q_norm = remove_accents(q_text)
    a_norm = remove_accents(full_a)
    
    # Tạo danh sách các lựa chọn đã chuẩn hóa
    options = []
    for opt in q_obj['raw_options']:
        options.append({
            'key': opt['key'],
            'text': opt['text'],
            'correct': opt['correct']
        })
    
    return {
        'id': qid,
        'q': q_text,
        'a': full_a,
        'ans_key': correct_key,
        'ans_text': correct_text,
        'options': options,
        'q_norm': q_norm,
        'a_norm': a_norm
    }

def main():
    excel_path = 'cauhoi.xlsx'
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found.")
        sys.exit(1)
        
    print(f"Reading {excel_path}...")
    questions = parse_excel(excel_path)
    print(f"Parsed {len(questions)} questions successfully.")
    
    # Ghi ra data.json
    with open('data.json', 'w', encoding='utf-8') as f:
        json.dump(questions, f, ensure_ascii=False, indent=2)
    print("Exported data.json")
    
    # Ghi ra data.js để có thể mở trực tiếp file:// mà không bị lỗi CORS browser
    with open('data.js', 'w', encoding='utf-8') as f:
        f.write("window.QUESTIONS_DATA = ")
        json.dump(questions, f, ensure_ascii=False, indent=2)
        f.write(";\n")
    print("Exported data.js")

if __name__ == '__main__':
    main()
